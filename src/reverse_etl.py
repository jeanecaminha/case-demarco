"""Reverse ETL: leva as recomendações do pipeline para a planilha dos gestores.

A regra principal é não estragar o histórico que os gestores escrevem à mão.
Por isso:
- as colunas são achadas pelo nome do cabeçalho, não pela posição;
- só as COLUNAS_CALCULADAS são escritas, célula a célula (nunca a planilha inteira);
- o cruzamento é pelo ID, porque a base tem homônimos;
- colaborador novo entra no fim; quem não está mais no pipeline fica como está;
- antes de gravar sai um backup, e depois a planilha é relida para conferir
  que as colunas manuais não mudaram.
"""
from __future__ import annotations

import csv
import logging
import os
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import pandas as pd

from .config import (
    COLUNA_CHAVE_PLANILHA,
    COLUNAS_CALCULADAS,
    COLUNAS_IDENTIFICACAO,
    COLUNAS_MANUAIS,
    PASTA_BACKUP,
    PASTA_SAIDA,
    Config,
)

log = logging.getLogger(__name__)


# O plano (o que vai mudar) é montado sem tocar em nada. É essa parte que os
# testes cobrem, e é ela que o --simular mostra.
@dataclass
class AtualizacaoCelula:
    linha: int          # número da linha na planilha (1 = cabeçalho)
    coluna: int         # número da coluna na planilha (1 = A)
    nome_coluna: str
    id_colaborador: str
    valor: object


@dataclass
class PlanoAtualizacao:
    atualizacoes: list[AtualizacaoCelula] = field(default_factory=list)
    novas_linhas: list[list[object]] = field(default_factory=list)
    ids_sem_recomendacao: list[str] = field(default_factory=list)

    def resumo(self) -> str:
        linhas_tocadas = len({a.linha for a in self.atualizacoes})
        return (
            f"{len(self.atualizacoes)} célula(s) em {linhas_tocadas} linha(s) existente(s); "
            f"{len(self.novas_linhas)} linha(s) nova(s); "
            f"{len(self.ids_sem_recomendacao)} linha(s) sem correspondência mantida(s) intacta(s)"
        )


def _normalizar_id(valor: object) -> str:
    return str(valor).strip().upper() if valor is not None else ""


def _mapa_colunas(cabecalho: list[object]) -> dict[str, int]:
    # nome da coluna -> índice. Se faltar alguma coluna esperada, não grava nada.
    mapa = {str(nome).strip(): i for i, nome in enumerate(cabecalho) if nome is not None and str(nome).strip()}
    obrigatorias = [COLUNA_CHAVE_PLANILHA, *COLUNAS_CALCULADAS, *COLUNAS_MANUAIS]
    faltantes = [c for c in obrigatorias if c not in mapa]
    if faltantes:
        raise ValueError(f"Planilha sem as colunas obrigatórias: {faltantes}. Nada foi gravado.")
    return mapa


def _valor_python(valor: object) -> object:
    # openpyxl e a API do Sheets não aceitam tipos do numpy
    if valor is None or (isinstance(valor, float) and pd.isna(valor)):
        return ""
    if hasattr(valor, "item"):
        return valor.item()
    return valor


def planejar_atualizacao(
    cabecalho: list[object],
    linhas: list[list[object]],
    recomendacoes: pd.DataFrame,
    data_processamento: str,
) -> PlanoAtualizacao:
    mapa = _mapa_colunas(cabecalho)
    idx_chave = mapa[COLUNA_CHAVE_PLANILHA]

    recs = {
        _normalizar_id(r[COLUNA_CHAVE_PLANILHA]): r
        for r in recomendacoes.to_dict("records")
    }
    plano = PlanoAtualizacao()
    ids_na_planilha = set()

    for posicao, linha in enumerate(linhas):
        numero_linha = posicao + 2  # linha 1 é o cabeçalho
        id_planilha = _normalizar_id(linha[idx_chave] if idx_chave < len(linha) else None)
        if not id_planilha:
            continue
        ids_na_planilha.add(id_planilha)
        rec = recs.get(id_planilha)
        if rec is None:
            plano.ids_sem_recomendacao.append(id_planilha)
            continue
        valores = {**{c: rec[c] for c in COLUNAS_CALCULADAS if c in rec}, "Data_Ultimo_Processamento": data_processamento}
        for nome_coluna in COLUNAS_CALCULADAS:
            plano.atualizacoes.append(
                AtualizacaoCelula(
                    linha=numero_linha,
                    coluna=mapa[nome_coluna] + 1,
                    nome_coluna=nome_coluna,
                    id_colaborador=id_planilha,
                    valor=_valor_python(valores[nome_coluna]),
                )
            )

    # quem está no pipeline mas não na planilha vira linha nova, com as colunas manuais vazias
    largura = max(mapa.values()) + 1
    for id_rec, rec in sorted(recs.items()):
        if id_rec in ids_na_planilha:
            continue
        nova = [""] * largura
        for nome_coluna in COLUNAS_IDENTIFICACAO + COLUNAS_CALCULADAS:
            if nome_coluna in mapa:
                valor = data_processamento if nome_coluna == "Data_Ultimo_Processamento" else rec.get(nome_coluna)
                nova[mapa[nome_coluna]] = _valor_python(valor)
        plano.novas_linhas.append(nova)

    return plano


def fotografia_colunas_manuais(cabecalho: list[object], linhas: list[list[object]]) -> dict[str, tuple]:
    # ID -> valores das colunas manuais, como texto, para comparar antes e depois
    mapa = _mapa_colunas(cabecalho)
    foto = {}
    for linha in linhas:
        id_colab = _normalizar_id(linha[mapa[COLUNA_CHAVE_PLANILHA]] if mapa[COLUNA_CHAVE_PLANILHA] < len(linha) else None)
        if not id_colab:
            continue
        foto[id_colab] = tuple(
            "" if (i := mapa[c]) >= len(linha) or linha[i] is None else str(linha[i]).strip()
            for c in COLUNAS_MANUAIS
        )
    return foto


def verificar_preservacao(antes: dict[str, tuple], depois: dict[str, tuple]) -> None:
    divergencias = {k: (v, depois.get(k)) for k, v in antes.items() if depois.get(k) != v}
    if divergencias:
        amostra = list(divergencias.items())[:5]
        raise RuntimeError(
            f"Colunas manuais mudaram em {len(divergencias)} linha(s) depois da gravação: {amostra}. "
            "Restaure pelo backup em data/backups/."
        )
    log.info("Conferência OK: colunas manuais de %d colaborador(es) intactas.", len(antes))


class PlanilhaGoogleSheets:
    def __init__(self, config: Config):
        import gspread

        if not config.gsheets_spreadsheet_id:
            raise ValueError("Defina GSHEETS_SPREADSHEET_ID (variável de ambiente ou .env).")
        credenciais = os.getenv("GOOGLE_APPLICATION_CREDENTIALS")
        cliente = gspread.service_account(filename=credenciais) if credenciais else gspread.service_account()
        planilha = cliente.open_by_key(config.gsheets_spreadsheet_id)
        self.aba = planilha.worksheet(config.gsheets_worksheet) if config.gsheets_worksheet else planilha.sheet1
        self.descricao = f"Google Sheets {config.gsheets_spreadsheet_id} / aba '{self.aba.title}'"

    def ler(self) -> tuple[list[object], list[list[object]]]:
        valores = self.aba.get_all_values()
        return valores[0], valores[1:]

    def aplicar(self, plano: PlanoAtualizacao) -> None:
        from gspread.utils import rowcol_to_a1

        # uma chamada só, mas célula a célula: se um gestor estiver editando
        # outra coluna nesse momento, a edição dele não é sobrescrita
        if plano.atualizacoes:
            self.aba.batch_update(
                [{"range": rowcol_to_a1(a.linha, a.coluna), "values": [[a.valor]]} for a in plano.atualizacoes],
                value_input_option="RAW",
            )
        if plano.novas_linhas:
            self.aba.append_rows(plano.novas_linhas, value_input_option="RAW")


class PlanilhaExcel:
    # Modo local: aplica o mesmo plano numa cópia .xlsx, mantendo a formatação.

    def __init__(self, origem: Path, destino: Path):
        import openpyxl

        self._openpyxl = openpyxl
        # a partir da 2ª execução atualiza a cópia, como seria com a planilha de verdade
        self.origem = destino if destino.exists() else origem
        self.destino = destino
        self.descricao = f"Excel {destino}"

    def _aba(self, caminho: Path):
        return self._openpyxl.load_workbook(caminho).worksheets[0]

    def ler(self, caminho: Path | None = None) -> tuple[list[object], list[list[object]]]:
        valores = [list(r) for r in self._aba(caminho or self.origem).iter_rows(values_only=True)]
        return valores[0], valores[1:]

    def aplicar(self, plano: PlanoAtualizacao) -> None:
        livro = self._openpyxl.load_workbook(self.origem)
        aba = livro.worksheets[0]
        for a in plano.atualizacoes:
            aba.cell(row=a.linha, column=a.coluna, value=a.valor)
        for nova in plano.novas_linhas:
            aba.append(nova)
        self.destino.parent.mkdir(parents=True, exist_ok=True)
        livro.save(self.destino)


def _salvar_backup(cabecalho: list[object], linhas: list[list[object]]) -> Path:
    PASTA_BACKUP.mkdir(parents=True, exist_ok=True)
    caminho = PASTA_BACKUP / f"planilha_gestores_{datetime.now():%Y%m%d_%H%M%S}.csv"
    with caminho.open("w", newline="", encoding="utf-8-sig") as f:
        escritor = csv.writer(f)
        escritor.writerow(cabecalho)
        escritor.writerows(linhas)
    return caminho


def executar_reverse_etl(
    alvo: str,
    config: Config,
    recomendacoes: pd.DataFrame,
    data_processamento: str,
    simular: bool = False,
) -> PlanoAtualizacao:
    if alvo == "bigquery":
        destino = PlanilhaGoogleSheets(config)
    else:
        destino = PlanilhaExcel(
            origem=config.planilha_gestores_local,
            destino=PASTA_SAIDA / "template_acoes_disciplinares_atualizado.xlsx",
        )

    cabecalho, linhas = destino.ler()
    plano = planejar_atualizacao(cabecalho, linhas, recomendacoes, data_processamento)
    log.info("Reverse ETL em %s: %s", destino.descricao, plano.resumo())
    if plano.ids_sem_recomendacao:
        log.warning("IDs na planilha que não estão no pipeline (mantidos): %s", plano.ids_sem_recomendacao)

    if simular:
        log.info("Simulação: nada foi gravado.")
        return plano

    backup = _salvar_backup(cabecalho, linhas)
    log.info("Backup da planilha salvo em %s", backup)
    foto_antes = fotografia_colunas_manuais(cabecalho, linhas)

    destino.aplicar(plano)

    if isinstance(destino, PlanilhaExcel):
        cab_depois, linhas_depois = destino.ler(destino.destino)
    else:
        cab_depois, linhas_depois = destino.ler()
    verificar_preservacao(foto_antes, fotografia_colunas_manuais(cab_depois, linhas_depois))
    return plano
